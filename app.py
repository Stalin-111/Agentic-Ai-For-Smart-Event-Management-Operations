import csv
import io
import os
import secrets
import smtplib
from collections import Counter
from datetime import datetime
from email.message import EmailMessage
from functools import wraps
from pathlib import Path

import qrcode
import requests
from flask import Flask, flash, jsonify, redirect, render_template, request, send_from_directory, session, url_for
from dotenv import load_dotenv
from werkzeug.utils import secure_filename

from database import get_db, init_db

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")
QR_DIR = BASE_DIR / "static" / "qr"
UPLOAD_DIR = BASE_DIR / "uploads"
ALLOWED_EXTENSIONS = {"csv", "xlsx"}

app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "change-this-before-deployment")
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = os.getenv("COOKIE_SECURE", "false").lower() == "true"


def admin_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if not session.get("is_admin"):
            flash("Please sign in as an administrator to open that page.", "warning")
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)
    return wrapped_view


def platform_snapshot():
    """Event Intelligence Engine: combines live data into executive decision signals."""
    with get_db() as db:
        attendee_row = db.execute("SELECT COUNT(*) AS total, COALESCE(SUM(checked_in), 0) AS checked FROM attendees").fetchone()
        sessions = db.execute("SELECT title, capacity FROM sessions ORDER BY start_time").fetchall()
        venues = db.execute("SELECT COUNT(*) AS count FROM venues").fetchone()["count"]
        speakers = db.execute("SELECT COUNT(*) AS count FROM speakers").fetchone()["count"]
        notifications = db.execute("SELECT COUNT(*) AS count FROM notifications WHERE status LIKE 'Delivered%'").fetchone()["count"]
        agent_runs = db.execute("SELECT * FROM agent_runs ORDER BY created_at DESC LIMIT 8").fetchall()
        attendee_events = Counter(row["event"] for row in db.execute("SELECT event FROM attendees").fetchall())
    total, checked = attendee_row["total"], attendee_row["checked"]
    checkin_rate = round((checked / total * 100) if total else 0)
    capacity_alerts = []
    for item in sessions:
        demand = attendee_events[item["title"]]
        if demand > item["capacity"]:
            capacity_alerts.append(f"{item['title']} exceeds capacity by {demand - item['capacity']} attendee(s).")
    readiness = min(100, round(35 + min(checkin_rate, 30) + min(len(sessions) * 4, 15) + min(venues * 5, 10) + min(speakers * 3, 10)))
    return {
        "total": total, "checked": checked, "checkin_rate": checkin_rate, "sessions": len(sessions),
        "venues": venues, "speakers": speakers, "notifications": notifications, "readiness": readiness,
        "capacity_alerts": capacity_alerts, "agent_runs": agent_runs, "event_demand": attendee_events,
    }


def intelligence_recommendations(snapshot):
    recommendations = []
    if snapshot["total"] == 0:
        recommendations.append("Begin attendee registration or import a registration file to activate meaningful event intelligence.")
    elif snapshot["checkin_rate"] < 50:
        recommendations.append(f"Check-in is {snapshot['checkin_rate']}%. Send a targeted arrival reminder and prepare QR support at the entrance.")
    else:
        recommendations.append(f"Check-in performance is healthy at {snapshot['checkin_rate']}%. Continue monitoring arrivals before each session.")
    recommendations.extend(snapshot["capacity_alerts"] or ["No current session capacity conflicts were detected."])
    if snapshot["venues"] == 0:
        recommendations.append("Add venue options to let the Venue Agent recommend the best allocation.")
    if snapshot["speakers"] == 0:
        recommendations.append("Add speaker profiles before finalising the event schedule.")
    return recommendations[:6]


def intelligence_command_center(snapshot):
    """Build the decision signals used by the Event Intelligence command center."""
    signals = []
    total = snapshot["total"]
    checked = snapshot["checked"]
    if total == 0:
        signals.append({"severity": "info", "title": "Waiting for attendee data", "detail": "Import or register attendees to activate demand and attendance signals."})
    elif snapshot["checkin_rate"] < 50:
        signals.append({"severity": "critical", "title": "Low check-in conversion", "detail": f"Only {snapshot['checkin_rate']}% of {total} registered attendees have checked in."})
    elif snapshot["checkin_rate"] < 80:
        signals.append({"severity": "warning", "title": "Check-in needs attention", "detail": f"{checked} of {total} attendees have checked in ({snapshot['checkin_rate']}%)."})
    else:
        signals.append({"severity": "success", "title": "Attendance flow is healthy", "detail": f"{snapshot['checkin_rate']}% check-in conversion is currently on track."})
    for alert in snapshot["capacity_alerts"]:
        signals.append({"severity": "critical", "title": "Session capacity risk", "detail": alert})
    if snapshot["venues"] == 0:
        signals.append({"severity": "warning", "title": "Venue coverage missing", "detail": "No venue options are available for allocation or optimisation."})
    if snapshot["speakers"] == 0:
        signals.append({"severity": "warning", "title": "Speaker coverage missing", "detail": "No speaker profiles are available for schedule validation."})
    if snapshot["sessions"] == 0:
        signals.append({"severity": "warning", "title": "Schedule is empty", "detail": "Create sessions to enable capacity and timetable intelligence."})
    if not signals:
        signals.append({"severity": "success", "title": "No active risks", "detail": "The command center has no unresolved operational signals."})
    severity_rank = {"critical": 0, "warning": 1, "info": 2, "success": 3}
    signals.sort(key=lambda item: severity_rank[item["severity"]])
    return signals[:8]


def run_agent_orchestration():
    """Coordinate the specialised agents in one auditable platform run."""
    snapshot = platform_snapshot()
    agent_results = [
        ("Registration Agent", "Registration analysis", "Healthy", f"Analysed {snapshot['total']} registrations and {snapshot['checked']} check-ins."),
        ("Venue Agent", "Venue optimisation", "Ready" if snapshot["venues"] else "Needs data", f"Reviewed {snapshot['venues']} venue option(s) against current demand."),
        ("Scheduling Agent", "Schedule validation", "Attention" if snapshot["capacity_alerts"] else "Healthy", " ".join(snapshot["capacity_alerts"]) or f"Reviewed {snapshot['sessions']} session(s); no capacity conflict detected."),
        ("Communications Agent", "Attendee communications", "Ready", f"{snapshot['notifications']} notification campaign(s) have been delivered."),
        ("Intelligence Agent", "Executive recommendations", "Healthy", intelligence_recommendations(snapshot)[0]),
    ]
    with get_db() as db:
        for agent, run_type, status, summary in agent_results:
            db.execute("INSERT INTO agent_runs (agent_name, run_type, status, summary) VALUES (?, ?, ?, ?)",
                       (agent, run_type, status, summary))
    return agent_results


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def make_qr(token):
    QR_DIR.mkdir(parents=True, exist_ok=True)
    filename = f"{token}.png"
    qrcode.make(token).save(QR_DIR / filename)
    return filename


def add_attendee(data):
    fields = ["name", "email", "phone", "age", "gender", "organization", "city", "category", "event"]
    cleaned = {key: str(data.get(key, "")).strip() for key in fields}
    if not all(cleaned.values()):
        raise ValueError("Please complete every registration field.")
    try:
        cleaned["age"] = int(cleaned["age"])
        if not 1 <= cleaned["age"] <= 120:
            raise ValueError
    except ValueError:
        raise ValueError("Age must be a number between 1 and 120.")

    token = secrets.token_urlsafe(12)
    with get_db() as db:
        cursor = db.execute("""
            INSERT INTO attendees (name, email, phone, age, gender, organization, city, category, event, qr_token)
            VALUES (:name, :email, :phone, :age, :gender, :organization, :city, :category, :event, :qr_token)
        """, {**cleaned, "qr_token": token})
        attendee_id = cursor.lastrowid
    make_qr(token)
    return attendee_id


def rows_to_chart_data(rows):
    def counts(column):
        result = Counter(row[column] for row in rows)
        return {"labels": list(result.keys()), "values": list(result.values())}

    age_groups = Counter()
    for row in rows:
        age = row["age"]
        label = "Under 18" if age < 18 else "18–24" if age <= 24 else "25–34" if age <= 34 else "35–44" if age <= 44 else "45+"
        age_groups[label] += 1
    daily = Counter(row["registration_time"][:10] for row in rows)
    return {
        "gender": counts("gender"), "category": counts("category"), "city": counts("city"),
        "event": counts("event"), "age": {"labels": list(age_groups.keys()), "values": list(age_groups.values())},
        "daily": {"labels": sorted(daily.keys()), "values": [daily[key] for key in sorted(daily)]},
    }


def build_insights(rows):
    if not rows:
        return ["No registrations yet. Add attendees to generate useful insights."]
    checked_in = sum(row["checked_in"] for row in rows)
    cities = Counter(row["city"] for row in rows)
    categories = Counter(row["category"] for row in rows)
    events = Counter(row["event"] for row in rows)
    ages = [row["age"] for row in rows]
    attendance = round(checked_in / len(rows) * 100)
    return [
        f"{categories.most_common(1)[0][0]} attendees form the largest audience segment.",
        f"{cities.most_common(1)[0][0]} has the highest number of registrations.",
        f"{events.most_common(1)[0][0]} is currently the most popular event.",
        f"Average attendee age is {sum(ages) / len(ages):.1f} years.",
        f"Current check-in rate is {attendance}%. Plan reminders for the {len(rows) - checked_in} pending attendee(s).",
    ]


def get_venue_recommendations(attendee_count, venues, requested_capacity=None, requested_city="", required_amenities=""):
    projected = max(1, requested_capacity or round(attendee_count * 1.15))
    amenities_needed = [item.strip().lower() for item in required_amenities.split(",") if item.strip()]
    recommendations = []
    for venue in venues:
        utilization = round(attendee_count / venue["capacity"] * 100) if venue["capacity"] else 0
        if venue["capacity"] >= projected:
            venue_amenities = (venue["amenities"] or "").lower()
            amenity_matches = sum(item in venue_amenities for item in amenities_needed)
            city_match = not requested_city or venue["city"].lower() == requested_city.lower()
            score = (venue["capacity"] - projected) + int(venue["cost"] / 1000) - (amenity_matches * 40) - (30 if city_match else 0)
            recommendations.append({"venue": venue, "utilization": utilization, "score": score,
                                    "amenity_matches": amenity_matches, "city_match": city_match})
    return sorted(recommendations, key=lambda item: item["score"]), projected


def notification_recipients(audience):
    with get_db() as db:
        if audience == "checked_in":
            rows = db.execute("SELECT email FROM attendees WHERE checked_in = 1").fetchall()
        elif audience == "pending":
            rows = db.execute("SELECT email FROM attendees WHERE checked_in = 0").fetchall()
        else:
            rows = db.execute("SELECT email FROM attendees").fetchall()
    return [row["email"] for row in rows]


def deliver_email_alert(recipients, subject, message):
    """Deliver an event alert through a configured SMTP account."""
    if not recipients:
        return "No recipient email addresses are available for the selected audience."
    host = os.getenv("SMTP_HOST")
    sender = os.getenv("SMTP_SENDER")
    if not host or not sender:
        return "Draft saved - configure SMTP_HOST and SMTP_SENDER to deliver emails."
    port = int(os.getenv("SMTP_PORT", "587"))
    username = os.getenv("SMTP_USERNAME", sender)
    password = os.getenv("SMTP_PASSWORD")
    if not password:
        return "Draft saved - configure SMTP_PASSWORD to deliver emails."
    use_ssl = os.getenv("SMTP_USE_SSL", "false").lower() in {"1", "true", "yes"}
    sent = 0
    smtp_class = smtplib.SMTP_SSL if use_ssl else smtplib.SMTP
    with smtp_class(host, port, timeout=20) as smtp:
        if not use_ssl:
            smtp.ehlo()
            smtp.starttls()
            smtp.ehlo()
        smtp.login(username, password)
        for recipient in recipients:
            email = EmailMessage()
            email["Subject"] = subject
            email["From"] = sender
            email["To"] = recipient
            email.set_content(message)
            smtp.send_message(email)
            sent += 1
    return f"Delivered to {sent} attendee(s)."


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("is_admin"):
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        expected_username = os.getenv("ADMIN_USERNAME", "admin")
        expected_password = os.getenv("ADMIN_PASSWORD", "EventFlow@2026")
        if secrets.compare_digest(username, expected_username) and secrets.compare_digest(password, expected_password):
            session.clear()
            session["is_admin"] = True
            flash("Welcome back, Administrator.", "success")
            return redirect(request.args.get("next") or url_for("dashboard"))
        flash("The username or password is incorrect.", "danger")
    return render_template("login.html")


@app.get("/logout")
def logout():
    session.clear()
    flash("You have been signed out.", "success")
    return redirect(url_for("index"))


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        try:
            attendee_id = add_attendee(request.form)
            flash("Registration saved and QR code created.", "success")
            return redirect(url_for("attendee_detail", attendee_id=attendee_id))
        except Exception as error:
            flash(str(error) if isinstance(error, ValueError) else "This email is already registered.", "danger")
    return render_template("register.html")


@app.route("/attendees")
@admin_required
def attendees():
    query = request.args.get("q", "").strip()
    with get_db() as db:
        rows = db.execute("SELECT * FROM attendees WHERE name LIKE ? OR email LIKE ? OR phone LIKE ? ORDER BY id DESC", [f"%{query}%"] * 3).fetchall()
    return render_template("attendees.html", attendees=rows, query=query)


@app.route("/attendee/<int:attendee_id>")
def attendee_detail(attendee_id):
    with get_db() as db:
        attendee = db.execute("SELECT * FROM attendees WHERE id = ?", (attendee_id,)).fetchone()
    if not attendee:
        flash("Attendee not found.", "danger")
        return redirect(url_for("attendees"))
    return render_template("attendee_detail.html", attendee=attendee)


@app.route("/dashboard")
@admin_required
def dashboard():
    with get_db() as db:
        rows = db.execute("SELECT * FROM attendees ORDER BY registration_time").fetchall()
    total = len(rows)
    checked_in = sum(row["checked_in"] for row in rows)
    today = datetime.now().strftime("%Y-%m-%d")
    today_count = sum(row["registration_time"].startswith(today) for row in rows)
    stats = {"total": total, "checked_in": checked_in, "pending": total - checked_in, "today": today_count,
             "average_age": round(sum((row["age"] for row in rows), 0) / total, 1) if total else 0,
             "cities": len(set(row["city"] for row in rows))}
    return render_template("dashboard.html", stats=stats, chart_data=rows_to_chart_data(rows))


@app.route("/analytics")
@admin_required
def analytics():
    with get_db() as db:
        rows = db.execute("SELECT * FROM attendees ORDER BY registration_time").fetchall()
    return render_template("analytics.html", chart_data=rows_to_chart_data(rows), total=len(rows))


@app.route("/checkin", methods=["GET", "POST"])
@admin_required
def checkin():
    attendee = None
    if request.method == "POST":
        lookup = request.form.get("lookup", "").strip()
        with get_db() as db:
            attendee = db.execute("SELECT * FROM attendees WHERE email = ? OR phone = ? OR qr_token = ?", (lookup, lookup, lookup)).fetchone()
        if not attendee:
            flash("No attendee matches that email, phone number, or QR token.", "danger")
    return render_template("checkin.html", attendee=attendee)


@app.post("/checkin/<int:attendee_id>/confirm")
@admin_required
def confirm_checkin(attendee_id):
    with get_db() as db:
        db.execute("UPDATE attendees SET checked_in = 1, checkin_time = COALESCE(checkin_time, CURRENT_TIMESTAMP) WHERE id = ?", (attendee_id,))
    flash("Check-in recorded successfully.", "success")
    return redirect(url_for("checkin"))


@app.route("/import", methods=["GET", "POST"])
@admin_required
def import_file():
    if request.method == "POST":
        upload = request.files.get("file")
        if not upload or not upload.filename or not allowed_file(upload.filename):
            flash("Upload a CSV or Excel (.xlsx) file.", "danger")
            return redirect(url_for("import_file"))
        filename = secure_filename(upload.filename)
        UPLOAD_DIR.mkdir(exist_ok=True)
        path = UPLOAD_DIR / filename
        upload.save(path)
        try:
            if filename.lower().endswith(".csv"):
                records = list(csv.DictReader(io.StringIO(path.read_text(encoding="utf-8-sig"))))
            else:
                from openpyxl import load_workbook
                sheet = load_workbook(path, read_only=True, data_only=True).active
                headers = [str(cell.value or "").strip().lower() for cell in next(sheet.iter_rows())]
                records = [dict(zip(headers, [cell.value or "" for cell in row])) for row in sheet.iter_rows(min_row=2)]
            added, skipped = 0, 0
            for record in records:
                try:
                    add_attendee({key.lower().strip(): value for key, value in record.items()})
                    added += 1
                except Exception:
                    skipped += 1
            flash(f"Import complete: {added} added, {skipped} skipped (duplicates or incomplete rows).", "success")
        except Exception as error:
            flash(f"Could not read file: {error}", "danger")
        return redirect(url_for("import_file"))
    return render_template("import.html")


@app.route("/insights")
@admin_required
def insights():
    with get_db() as db:
        rows = db.execute("SELECT * FROM attendees").fetchall()
    calculated = build_insights(rows)
    ai_text = None
    api_key = os.getenv("GEMINI_API_KEY")
    if request.args.get("ai") == "1" and api_key and rows:
        prompt = "Give 4 concise event-organizer recommendations from this attendee summary: " + "; ".join(calculated)
        try:
            response = requests.post(
                f"https://generativelanguage.googleapis.com/v1beta/models/{os.getenv('GEMINI_MODEL', 'gemini-2.0-flash')}:generateContent?key={api_key}",
                json={"contents": [{"parts": [{"text": prompt}]}]}, timeout=20,
            )
            response.raise_for_status()
            ai_text = response.json()["candidates"][0]["content"]["parts"][0]["text"]
        except Exception:
            flash("Gemini could not be reached, so calculated insights are shown instead.", "warning")
    return render_template("insights.html", insights=calculated, ai_text=ai_text, ai_enabled=bool(api_key))


@app.route("/notifications", methods=["GET", "POST"])
@admin_required
def notifications():
    if request.method == "POST":
        action = request.form.get("action", "send")
        subject = request.form.get("subject", "").strip()
        message = request.form.get("message", "").strip()
        audience = request.form.get("audience", "all")
        if action == "test":
            test_email = request.form.get("test_email", "").strip()
            if not test_email:
                flash("Enter an email address for the delivery test.", "danger")
                return redirect(url_for("notifications"))
            try:
                status = deliver_email_alert([test_email], "EventFlow email delivery test",
                                             "Your EventFlow notification service is configured and ready to send event alerts.")
                category = "success" if status.startswith("Delivered") else "warning"
            except Exception:
                status = "Delivery test failed. Confirm your SMTP host, port, username, and password."
                category = "danger"
            with get_db() as db:
                db.execute("INSERT INTO notifications (subject, message, audience, recipient_count, status) VALUES (?, ?, ?, ?, ?)",
                           ("EventFlow email delivery test", "SMTP delivery test", "test", 1, status))
            flash(status, category)
            return redirect(url_for("notifications"))
        if not subject or not message:
            flash("Provide both a subject and message for the notification.", "danger")
        else:
            recipients = notification_recipients(audience)
            try:
                status = deliver_email_alert(recipients, subject, message)
                category = "success" if status.startswith("Delivered") else "warning"
            except Exception as error:
                status = f"Delivery failed: {error}"
                category = "danger"
            with get_db() as db:
                db.execute("INSERT INTO notifications (subject, message, audience, recipient_count, status) VALUES (?, ?, ?, ?, ?)",
                           (subject, message, audience, len(recipients), status))
            flash(status, category)
            return redirect(url_for("notifications"))
    with get_db() as db:
        history = db.execute("SELECT * FROM notifications ORDER BY created_at DESC LIMIT 10").fetchall()
        total = db.execute("SELECT COUNT(*) AS count FROM attendees").fetchone()["count"]
        checked = db.execute("SELECT COUNT(*) AS count FROM attendees WHERE checked_in = 1").fetchone()["count"]
    smtp_ready = all(os.getenv(key) for key in ("SMTP_HOST", "SMTP_SENDER", "SMTP_PASSWORD"))
    return render_template("notifications.html", history=history, total=total, checked=checked, smtp_ready=smtp_ready)


@app.route("/venues", methods=["GET", "POST"])
@admin_required
def venues():
    if request.method == "POST":
        try:
            name = request.form.get("name", "").strip()
            city = request.form.get("city", "").strip()
            capacity = int(request.form.get("capacity", 0))
            cost = float(request.form.get("cost", 0))
            amenities = request.form.get("amenities", "").strip()
            if not name or not city or capacity < 1 or cost < 0:
                raise ValueError
            with get_db() as db:
                db.execute("INSERT INTO venues (name, city, capacity, cost, amenities) VALUES (?, ?, ?, ?, ?)",
                           (name, city, capacity, cost, amenities))
            flash("Venue added to the planning library.", "success")
            return redirect(url_for("venues"))
        except ValueError:
            flash("Enter a venue name, city, valid capacity, and non-negative cost.", "danger")
    with get_db() as db:
        venue_rows = db.execute("SELECT * FROM venues ORDER BY capacity ASC").fetchall()
        attendee_count = db.execute("SELECT COUNT(*) AS count FROM attendees").fetchone()["count"]
    try:
        requested_capacity = int(request.args.get("required_capacity", "0")) or None
    except ValueError:
        requested_capacity = None
    requested_city = request.args.get("city", "").strip()
    required_amenities = request.args.get("amenities", "").strip()
    recommendations, projected = get_venue_recommendations(attendee_count, venue_rows, requested_capacity, requested_city, required_amenities)
    workflows = [
        f"Plan for {projected} seats: registrations plus a 15% operational buffer.",
        f"Deploy {max(2, round(projected / 75))} check-in staff member(s) for the expected arrival volume.",
        f"Use {max(1, round(projected / 150))} entry lane(s), with a separate QR support desk.",
        "Review venue amenities, accessibility, Wi-Fi, AV, and emergency access before confirmation.",
    ]
    return render_template("venues.html", venues=venue_rows, recommendations=recommendations, attendee_count=attendee_count,
                           projected=projected, workflows=workflows, requested_capacity=requested_capacity or "",
                           requested_city=requested_city, required_amenities=required_amenities)


@app.route("/speakers", methods=["GET", "POST"])
@admin_required
def speakers():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        email = request.form.get("email", "").strip()
        availability = request.form.get("availability", "").strip()
        if not name or not email or not availability:
            flash("Name, email, and availability are required for every speaker profile.", "danger")
        else:
            try:
                with get_db() as db:
                    db.execute("""INSERT INTO speakers (name, email, organization, expertise, availability, bio)
                        VALUES (?, ?, ?, ?, ?, ?)""", (name, email, request.form.get("organization", "").strip(),
                        request.form.get("expertise", "").strip(), availability, request.form.get("bio", "").strip()))
                flash("Speaker profile added successfully.", "success")
                return redirect(url_for("speakers"))
            except Exception:
                flash("A speaker profile already exists for that email address.", "danger")
    with get_db() as db:
        speaker_rows = db.execute("SELECT * FROM speakers ORDER BY name").fetchall()
        assignments = db.execute("SELECT speaker_id, COUNT(*) AS count FROM sessions WHERE speaker_id IS NOT NULL GROUP BY speaker_id").fetchall()
    assignment_counts = {row["speaker_id"]: row["count"] for row in assignments}
    return render_template("speakers.html", speakers=speaker_rows, assignment_counts=assignment_counts)


@app.route("/schedule", methods=["GET", "POST"])
@admin_required
def schedule():
    with get_db() as db:
        venues = db.execute("SELECT * FROM venues ORDER BY name").fetchall()
        speaker_rows = db.execute("SELECT * FROM speakers ORDER BY name").fetchall()
    if request.method == "POST":
        try:
            title = request.form.get("title", "").strip()
            speaker_id = int(request.form.get("speaker_id", 0))
            speaker = next((row for row in speaker_rows if row["id"] == speaker_id), None)
            speaker_name = speaker["name"] if speaker else ""
            speaker_email = speaker["email"] if speaker else ""
            venue_id = int(request.form.get("venue_id", 0))
            start_time = request.form.get("start_time", "")
            end_time = request.form.get("end_time", "")
            capacity = int(request.form.get("capacity", 0))
            if not title or not speaker_name or not venue_id or not start_time or not end_time or capacity < 1 or start_time >= end_time:
                raise ValueError
            with get_db() as db:
                conflict = db.execute("""SELECT id FROM sessions
                    WHERE (venue_id = ? OR lower(speaker_name) = lower(?))
                    AND start_time < ? AND end_time > ?""", (venue_id, speaker_name, end_time, start_time)).fetchone()
                if conflict:
                    raise RuntimeError("Schedule conflict: the selected venue or speaker is already booked at this time.")
                db.execute("""INSERT INTO sessions (title, speaker_name, speaker_email, speaker_id, venue_id, start_time, end_time, capacity)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)""", (title, speaker_name, speaker_email, speaker_id, venue_id, start_time, end_time, capacity))
            flash("Session scheduled successfully.", "success")
            return redirect(url_for("schedule"))
        except RuntimeError as error:
            flash(str(error), "danger")
        except (TypeError, ValueError):
            flash("Complete all required session fields and ensure the end time is later than the start time.", "danger")
    with get_db() as db:
        sessions = db.execute("""SELECT sessions.*, venues.name AS venue_name FROM sessions
            JOIN venues ON venues.id = sessions.venue_id ORDER BY sessions.start_time""").fetchall()
    return render_template("schedule.html", venues=venues, speakers=speaker_rows, sessions=sessions)


@app.route("/session-analytics")
@admin_required
def session_analytics():
    with get_db() as db:
        sessions = db.execute("""SELECT sessions.*, venues.name AS venue_name FROM sessions
            JOIN venues ON venues.id = sessions.venue_id ORDER BY sessions.start_time""").fetchall()
        attendees = db.execute("SELECT event, checked_in FROM attendees").fetchall()
    attendance = Counter(row["event"] for row in attendees)
    checked_in = Counter(row["event"] for row in attendees if row["checked_in"])
    session_data = []
    for session in sessions:
        registered = attendance[session["title"]]
        arrived = checked_in[session["title"]]
        session_data.append({"session": session, "registered": registered, "arrived": arrived,
                             "fill_rate": round(registered / session["capacity"] * 100) if session["capacity"] else 0})
    stats = {"sessions": len(sessions), "capacity": sum(row["capacity"] for row in sessions),
             "registrations": sum(attendance.values()), "checked_in": sum(checked_in.values())}
    chart_data = {"labels": [item["session"]["title"] for item in session_data],
                  "registered": [item["registered"] for item in session_data],
                  "arrived": [item["arrived"] for item in session_data]}
    return render_template("session_analytics.html", sessions=session_data, stats=stats, chart_data=chart_data)


@app.route("/intelligence", methods=["GET", "POST"])
@admin_required
def intelligence():
    if request.method == "POST":
        results = run_agent_orchestration()
        attention = sum(result[2] in {"Attention", "Needs data"} for result in results)
        flash(f"Agent orchestration completed. {len(results)} agents ran; {attention} item(s) need attention.", "success")
        return redirect(url_for("intelligence"))
    snapshot = platform_snapshot()
    return render_template("intelligence.html", snapshot=snapshot, recommendations=intelligence_recommendations(snapshot), signals=intelligence_command_center(snapshot))


@app.route("/executive")
@admin_required
def executive_dashboard():
    snapshot = platform_snapshot()
    return render_template("executive.html", snapshot=snapshot, recommendations=intelligence_recommendations(snapshot))


@app.route("/platform-health")
@admin_required
def platform_health():
    database_ok = True
    try:
        with get_db() as db:
            db.execute("SELECT 1").fetchone()
    except Exception:
        database_ok = False
    health = [
        {"service": "Application server", "status": "Healthy", "detail": "Flask application is responding."},
        {"service": "SQLite data store", "status": "Healthy" if database_ok else "Unavailable", "detail": "Database connection check completed." if database_ok else "Database connection could not be verified."},
        {"service": "Session security", "status": "Healthy", "detail": "Administrator routes require an authenticated session."},
        {"service": "Email notifications", "status": "Configured" if all(os.getenv(key) for key in ("SMTP_HOST", "SMTP_SENDER", "SMTP_PASSWORD")) else "Optional setup", "detail": "SMTP delivery is available when provider credentials are configured."},
    ]
    return render_template("platform_health.html", health=health, snapshot=platform_snapshot())


@app.get("/api/intelligence/status")
@admin_required
def intelligence_status_api():
    snapshot = platform_snapshot()
    return jsonify({
        "readiness": snapshot["readiness"], "checkin_rate": snapshot["checkin_rate"],
        "registrations": snapshot["total"], "checked_in": snapshot["checked"],
        "capacity_alerts": snapshot["capacity_alerts"], "recommendations": intelligence_recommendations(snapshot),
        "signals": intelligence_command_center(snapshot),
    })


@app.route("/qr/<path:filename>")
def qr_file(filename):
    return send_from_directory(QR_DIR, filename)


if __name__ == "__main__":
    init_db()
    app.run(host="127.0.0.1", debug=os.getenv("FLASK_DEBUG", "false").lower() == "true")
